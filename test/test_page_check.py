"""page_check: source text vs docling's page markdown (synthetic text only)."""

import json

import pytest

import artmind.ingest as ing
from artmind import page_check
from artmind.page_check import (
    PageFinding,
    check_pages,
    coverage,
    page_repair_mode,
    source_pages,
    strip_furniture,
    structural_flags,
)

SENTENCE = "The quick brown fox jumps over the lazy dog near the old river bank"


# ── mode ──────────────────────────────────────────────────────────────────────


@pytest.mark.parametrize(
    "value, expected",
    [(None, "detect"), ("", "detect"), ("off", "off"), ("DETECT", "detect"), (" repair ", "repair"), ("bogus", "detect")],
)
def test_mode_defaults_to_detect(value, expected):
    env = {} if value is None else {"ARTMIND_PAGE_REPAIR": value}
    assert page_repair_mode(env) == expected


# ── coverage ──────────────────────────────────────────────────────────────────


def test_identical_text_is_fully_covered():
    cov, missing = coverage(SENTENCE, SENTENCE)
    assert cov == 1.0 and missing == []


def test_a_lost_last_line_flags_the_page():
    source = f"{SENTENCE}\nCalculation:\nvalue equals total divided by count"
    markdown = f"{SENTENCE}\n\nCalculation:"
    cov, missing = coverage(source, markdown)
    assert cov < 0.97
    assert any("divided" in m for m in missing)
    [finding] = check_pages([source], {1: markdown})
    assert (finding.page, finding.reason) == (1, "missing-text") and finding.coverage < 0.97


def test_entities_curly_quotes_and_nbsp_do_not_flag():
    source = "Terms & conditions apply to the client’s “standard” offer\xa0only today"
    markdown = "Terms &amp; conditions apply to the client's \"standard\" offer only today"
    assert coverage(source, markdown)[0] == 1.0


def test_reading_order_across_layout_fragments_creates_no_phantom_phrases():
    # two table columns on one line: the markdown reads them in another order
    source = "alpha beta gamma delta     one two three four"
    markdown = "one two three four\n\nalpha beta gamma delta"
    assert coverage(source, markdown)[0] == 1.0


def test_markup_in_the_markdown_does_not_matter():
    source = "Heading words here\nfirst point is this\nsecond point is that"
    markdown = "## Heading words here\n\n- first point is this\n- second point is that\n"
    assert coverage(source, markdown)[0] == 1.0


def test_a_hyphenation_mark_in_the_text_layer_is_not_part_of_the_word():
    source = "the visitor was neither ­willing nor ei￾ther able to stay"
    markdown = "the visitor was neither willing nor either able to stay"
    assert coverage(source, markdown)[0] == 1.0


def test_words_docling_split_or_spaced_out_still_match():
    source = "every possible combination of the 6.30 train times"
    markdown = "every possible combi nation of the 6 . 30 train times"
    assert coverage(source, markdown)[0] == 1.0


def test_a_phrase_the_markdown_lacks_is_still_missing_after_squashing():
    cov, missing = coverage("alpha beta gamma delta epsilon", "alpha beta gamma delta")
    assert cov < 1.0 and missing == ["gamma delta epsilon"]


def test_a_paragraph_that_runs_over_a_page_break_is_not_a_loss():
    """docling keeps the whole paragraph on the page it starts on."""
    first = "the opening of a long paragraph that carries on over the page break"
    second = "and finishes here with the rest of the closing sentence"
    source = [f"{first}\n{second[:20]}", f"{second[20:]}\nanother line of that page here"]
    markdown = {1: f"{first} {second}", 2: "another line of that page here"}
    assert check_pages(source, markdown) == []


def test_text_missing_from_every_neighbouring_page_is_still_flagged():
    source = ["alpha beta gamma delta", "epsilon zeta eta theta", "iota kappa lambda mu"]
    markdown = {1: "alpha beta gamma delta", 2: "epsilon zeta", 3: "iota kappa lambda mu"}
    assert [f.page for f in check_pages(source, markdown)] == [2]


def test_a_page_with_nothing_to_compare_is_covered():
    assert coverage("", "anything at all here") == (1.0, [])
    assert coverage("a b", "unrelated") == (1.0, [])  # fewer words than a phrase


def test_a_clean_page_is_not_flagged():
    assert check_pages([SENTENCE], {1: SENTENCE}) == []


# ── page furniture ────────────────────────────────────────────────────────────


def _pages(n, body):
    return [f"Running header text\n{body(i)}\nPage {i} of {n}" for i in range(1, n + 1)]


def test_header_and_footer_repeated_on_most_pages_are_dropped():
    words = ["apple", "birch", "cedar", "dune", "ember", "fjord", "grove", "heath", "iris", "jade"]
    pages = _pages(10, lambda i: f"body line about {words[i - 1]}")
    stripped = strip_furniture(pages)
    assert all("Running header" not in p and "Page" not in p for p in stripped)
    assert all(f"body line about {words[i]}" in p for i, p in enumerate(stripped))


def test_furniture_the_markdown_omits_does_not_flag_the_page():
    pages = _pages(10, lambda i: f"{SENTENCE} number {i}")
    markdown = {i: f"{SENTENCE} number {i}" for i in range(1, 11)}
    assert check_pages(pages, markdown) == []


def test_a_body_paragraph_repeated_on_three_pages_is_kept():
    """The prototype's regression: stripping any repeated line removed real
    content. Only a page's first/last three lines are candidates."""
    paragraph = "This paragraph repeats verbatim in the middle of three pages"
    pages = []
    for i in range(1, 11):
        filler = [f"filler line {i} number {j} here" for j in range(1, 8)]
        if i in (2, 5, 8):
            filler.insert(3, paragraph)
        pages.append("\n".join(filler))
    stripped = strip_furniture(pages)
    for i in (2, 5, 8):
        assert paragraph in stripped[i - 1]


def test_a_short_document_keeps_its_edge_lines():
    pages = [f"Only page {i} line one\nOnly page {i} line two" for i in (1, 2)]
    assert strip_furniture(pages) == pages


# ── structure ─────────────────────────────────────────────────────────────────


def test_a_long_cell_ending_in_a_colon_is_a_sign_of_loss():
    cell = "Calculation of the adjusted figure for the period under review is defined as:"
    assert structural_flags(f"| a | {cell} |\n|---|---|\n") == ["cell-ends-with-colon"]
    assert structural_flags("| a | Short: |\n") == []


def test_long_repeated_rows_are_a_sign_of_loss():
    row = "| " + " | ".join(["word " * 12] * 4) + " |"
    assert structural_flags(f"{row}\n{row}\n") == ["repeated-rows"]
    assert structural_flags(f"{row}\n") == []


def test_a_structural_sign_flags_a_page_the_text_check_passes():
    cell = "Calculation of the adjusted figure for the period under review is defined as:"
    md = f"| name | {cell} |\n|---|---|\n"
    [finding] = check_pages([f"name {cell}"], {1: md})
    assert finding.reason == "cell-ends-with-colon"


def test_a_table_split_across_pages_is_covered_page_by_page():
    source = ["row one cell text    row one other\nrow two cell text    row two other", "row three cell text    row three other"]
    markdown = {
        1: "| row one cell text | row one other |\n| row two cell text | row two other |",
        2: "| row three cell text | row three other |",
    }
    assert check_pages(source, markdown) == []


def test_the_record_holds_no_document_text():
    finding = PageFinding(3, 0.8234, "missing-text", ["secret words here"])
    assert finding.record() == {"page": 3, "coverage": 0.823, "reason": "missing-text"}


# ── source text ───────────────────────────────────────────────────────────────


def _pdf_bytes(pages: list[str]) -> bytes:
    """A minimal text-layer PDF, one line of Helvetica per page ('' = no text)."""
    objects = ["<< /Type /Catalog /Pages 2 0 R >>"]
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(len(pages)))
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>")
    font = 3 + 2 * len(pages)
    for i, text in enumerate(pages):
        content = f"BT /F1 12 Tf 50 700 Td ({text}) Tj ET" if text else ""
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {4 + 2 * i} 0 R "
            f"/Resources << /Font << /F1 {font} 0 R >> >> >>"
        )
        objects.append(f"<< /Length {len(content)} >>\nstream\n{content}\nendstream")
    objects.append("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out = b"%PDF-1.4\n"
    offsets = []
    for n, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{n} 0 obj\n{obj}\nendobj\n".encode()
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += "".join(f"{o:010d} 00000 n \n" for o in offsets).encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return out


def test_pdf_page_text_comes_from_the_text_layer(tmp_path):
    pdf = tmp_path / "doc.pdf"
    pdf.write_bytes(_pdf_bytes(["first page words", "second page words"]))
    pages = source_pages(pdf)
    assert len(pages) == 2
    assert "first page words" in pages[0] and "second page words" in pages[1]


def test_other_types_have_no_source_pages(tmp_path):
    docx = tmp_path / "doc.docx"
    docx.write_bytes(b"x")
    assert source_pages(docx) is None


def test_an_unreadable_pdf_has_no_pages_and_does_not_raise(tmp_path):
    pdf = tmp_path / "broken.pdf"
    pdf.write_bytes(b"not a pdf")
    assert source_pages(pdf) == []


def test_a_missing_pdf_library_skips_the_check_instead_of_failing(tmp_path, monkeypatch):
    import builtins

    real_import = builtins.__import__

    def no_pdfium(name, *args, **kwargs):
        if name == "pypdfium2":
            raise ImportError("No module named 'pypdfium2'")
        return real_import(name, *args, **kwargs)

    pdf = tmp_path / "doc.pdf"
    pdf.write_bytes(_pdf_bytes(["alpha beta gamma delta"]))
    monkeypatch.setattr(builtins, "__import__", no_pdfium)

    check = page_check.check_document(pdf, _docling_doc({1: "alpha beta gamma delta"}))

    assert check.status == "skipped" and check.findings == []


def _make_pptx(path, slides, notes=None):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    for i, lines in enumerate(slides, start=1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        box.text_frame.text = lines[0]
        for extra in lines[1:]:
            box.text_frame.add_paragraph().text = extra
        if i == 1:
            table = slide.shapes.add_table(1, 2, Inches(1), Inches(3), Inches(4), Inches(1)).table
            table.cell(0, 0).text = "table cell alpha"
            table.cell(0, 1).text = "table cell beta"
            group = slide.shapes.add_group_shape()
            inner = group.shapes.add_textbox(Inches(1), Inches(5), Inches(2), Inches(1))
            inner.text_frame.text = "grouped shape text"
        if notes and i in notes:
            slide.notes_slide.notes_text_frame.text = notes[i]
    prs.save(str(path))


def test_pptx_text_covers_shapes_tables_groups_and_notes(tmp_path):
    deck = tmp_path / "deck.pptx"
    _make_pptx(deck, [["slide one text", "second paragraph"], ["slide two text"]], notes={2: "note for two"})
    pages = source_pages(deck)
    assert len(pages) == 2
    for phrase in ("slide one text", "second paragraph", "table cell alpha", "table cell beta", "grouped shape text"):
        assert phrase in pages[0]
    assert "slide two text" in pages[1] and "note for two" in pages[1]


# ── a document ────────────────────────────────────────────────────────────────


def _docling_doc(pages: dict[int, str], notes: dict[int, str] | None = None):
    pytest.importorskip("docling_core")
    from docling_core.types.doc import (
        BoundingBox, ContentLayer, DocItemLabel, DoclingDocument, ProvenanceItem, Size,
    )

    def prov(n, text):
        return ProvenanceItem(page_no=n, bbox=BoundingBox(l=0, t=0, r=1, b=1), charspan=(0, len(text)))

    doc = DoclingDocument(name="d")
    for n, text in pages.items():
        doc.add_page(page_no=n, size=Size(width=100, height=100))
        doc.add_text(label=DocItemLabel.TEXT, text=text, prov=prov(n, text))
    for n, text in (notes or {}).items():
        doc.add_text(label=DocItemLabel.TEXT, text=text, prov=prov(n, text), content_layer=ContentLayer.NOTES)
    return doc


def test_check_document_flags_only_the_page_that_lost_text(tmp_path):
    pdf = tmp_path / "doc.pdf"
    pdf.write_bytes(_pdf_bytes(["alpha beta gamma delta epsilon", "zeta eta theta iota kappa"]))
    doc = _docling_doc({1: "alpha beta gamma delta epsilon", 2: "zeta eta"})

    check = page_check.check_document(pdf, doc)

    assert check.status == "checked" and check.pages == 2
    assert [f.page for f in check.findings] == [2]
    assert check.source[2] and check.markdown[1]


def test_a_pdf_without_a_text_layer_is_skipped_never_flagged(tmp_path):
    pdf = tmp_path / "scan.pdf"
    pdf.write_bytes(_pdf_bytes(["", ""]))
    doc = _docling_doc({1: "text docling read by ocr here", 2: "more text read by ocr here"})

    check = page_check.check_document(pdf, doc)

    assert check.status == "skipped" and check.why == "no text layer" and check.findings == []


def test_a_page_count_mismatch_is_skipped(tmp_path):
    pdf = tmp_path / "doc.pdf"
    pdf.write_bytes(_pdf_bytes(["alpha beta gamma delta"]))
    doc = _docling_doc({1: "alpha beta gamma delta", 2: "extra page words here"})
    assert page_check.check_document(pdf, doc).status == "skipped"


def test_other_types_are_skipped(tmp_path):
    docx = tmp_path / "doc.docx"
    docx.write_bytes(b"x")
    assert page_check.check_document(docx, _docling_doc({1: "words here now"})).status == "skipped"


def test_a_pptx_note_missing_from_the_markdown_is_a_finding(tmp_path):
    deck = tmp_path / "deck.pptx"
    _make_pptx(deck, [["slide one text"], ["slide two text"]], notes={2: "remember to mention the quarterly figures"})
    doc = _docling_doc({1: "slide one text table cell alpha table cell beta grouped shape text", 2: "slide two text"})

    check = page_check.check_document(deck, doc)

    assert [f.page for f in check.findings] == [2]


def test_a_pptx_with_its_notes_in_the_markdown_is_clean(tmp_path):
    deck = tmp_path / "deck.pptx"
    _make_pptx(deck, [["slide one text"], ["slide two text"]], notes={2: "remember to mention the quarterly figures"})
    doc = _docling_doc(
        {1: "slide one text table cell alpha table cell beta grouped shape text", 2: "slide two text"},
        notes={2: "remember to mention the quarterly figures"},
    )
    assert page_check.check_document(deck, doc).findings == []


def test_summary_counts():
    check = page_check.PageCheck("checked", pages=5, findings=[PageFinding(2, 0.5, "x"), PageFinding(4, 0.6, "x")])
    assert page_check.summary(check, repaired=[2]) == {"flagged": 2, "repaired": 1, "unrepaired": 1}


# ── wired into the conversion ────────────────────────────────────────────────


def _stub_docling_json(monkeypatch, doc, md):
    def run(cmd, *args, **kwargs):
        (ing.MARKDOWNS_DIR / "deck.md").write_text(md, encoding="utf-8")
        doc.save_as_json(ing.MARKDOWNS_DIR / "deck.json")
        return 0, "", ""

    monkeypatch.setattr(ing, "run_command", run)
    monkeypatch.setattr(ing, "_describe_image", lambda *a, **k: pytest.fail("no image expected"))


def _pdf_source(ingest_env):
    vault, source = ingest_env
    pdf = source.with_suffix(".pdf")
    pdf.write_bytes(_pdf_bytes(["alpha beta gamma delta epsilon", "zeta eta theta iota kappa"]))
    return pdf


def _sidecar():
    return json.loads((ing.MARKDOWNS_DIR / "deck_chunks" / ing.CONVERSION_SIDECAR).read_text())


def test_detect_records_findings_logs_one_warning_and_counts(ingest_env, monkeypatch):
    from loguru import logger

    pdf = _pdf_source(ingest_env)
    doc = _docling_doc({1: "alpha beta gamma delta epsilon", 2: "zeta eta"})
    _stub_docling_json(monkeypatch, doc, "alpha beta gamma delta epsilon\n\nzeta eta\n")
    monkeypatch.delenv("ARTMIND_PAGE_REPAIR", raising=False)
    warnings = []
    sink = logger.add(lambda m: warnings.append(m), level="WARNING", format="{message}")
    try:
        result = ing.ingest_file(pdf, "model", "general", chunk_size=6000)
    finally:
        logger.remove(sink)

    conversion = _sidecar()["conversion"]
    assert conversion["page_check"] == "detect" and conversion["pages"] == 2
    assert [f["page"] for f in conversion["flagged"]] == [2]
    assert conversion["flagged"][0]["reason"] == "missing-text"
    assert conversion["repaired"] == [] and conversion["unrepaired"] == [2]
    assert "iota" not in json.dumps(conversion)  # no document text in the record
    assert result["page_check"] == {"flagged": 1, "repaired": 0, "unrepaired": 1}
    flagged_lines = [w for w in warnings if "lack source text" in w]
    assert len(flagged_lines) == 1 and "2 (" in flagged_lines[0]


def test_off_skips_the_check(ingest_env, monkeypatch):
    pdf = _pdf_source(ingest_env)
    _stub_docling_json(monkeypatch, _docling_doc({1: "alpha beta gamma delta epsilon", 2: "zeta eta"}), "x y z\n")
    monkeypatch.setenv("ARTMIND_PAGE_REPAIR", "off")
    monkeypatch.setattr(page_check, "check_document", lambda *a: pytest.fail("must not run"))

    result = ing.ingest_file(pdf, "model", "general", chunk_size=6000)

    assert _sidecar()["conversion"]["page_check"] == "off"
    assert "page_check" not in result


def test_a_clean_pdf_records_no_flags(ingest_env, monkeypatch):
    pdf = _pdf_source(ingest_env)
    doc = _docling_doc({1: "alpha beta gamma delta epsilon", 2: "zeta eta theta iota kappa"})
    _stub_docling_json(monkeypatch, doc, "alpha beta gamma delta epsilon\n\nzeta eta theta iota kappa\n")
    monkeypatch.delenv("ARTMIND_PAGE_REPAIR", raising=False)

    result = ing.ingest_file(pdf, "model", "general", chunk_size=6000)

    conversion = _sidecar()["conversion"]
    assert conversion["flagged"] == [] and conversion["pages"] == 2
    assert "page_check" not in result


def test_a_missing_json_marks_the_check_skipped(ingest_env, monkeypatch):
    pdf = _pdf_source(ingest_env)

    def run(cmd, *args, **kwargs):
        (ing.MARKDOWNS_DIR / "deck.md").write_text("alpha beta\n", encoding="utf-8")
        return 0, "", ""

    monkeypatch.setattr(ing, "run_command", run)
    monkeypatch.delenv("ARTMIND_PAGE_REPAIR", raising=False)

    ing.ingest_file(pdf, "model", "general", chunk_size=6000)

    conversion = _sidecar()["conversion"]
    assert conversion["page_check"] == "skipped" and conversion["why"] == "docling JSON unusable"


def test_job_status_shows_the_counts(tmp_path, monkeypatch):
    import artmind.db as db
    from artmind import jobs

    monkeypatch.setattr(db, "DB_PATH", tmp_path / "registry.db")
    conn = db._get_db()
    conn.execute(
        "INSERT INTO ingestion_jobs (job_id, status, file_count, processed_count, queued_at) "
        "VALUES ('j1', 'completed', 1, 1, 'now')"
    )
    conn.execute("INSERT INTO ingestion_job_files (job_id, filename, status) VALUES ('j1', 'a.pdf', 'completed')")
    conn.commit()
    conn.close()

    jobs._update_job_file_status("j1", "a.pdf", page_check={"flagged": 2, "repaired": 0, "unrepaired": 2})

    entry = jobs._get_job_status("j1")["files"][0]
    assert entry["page_check"] == {"flagged": 2, "repaired": 0, "unrepaired": 2}


def test_text_docling_recorded_on_a_later_page_but_exports_earlier_is_not_a_loss(tmp_path):
    """An item whose provenance spans pages 1 and 3 exports on page 1 only;
    page 3's source text is then misplaced in the body, not missing."""
    pytest.importorskip("docling_core")
    from docling_core.types.doc import BoundingBox, DocItemLabel, DoclingDocument, ProvenanceItem, Size

    pdf = tmp_path / "doc.pdf"
    pdf.write_bytes(_pdf_bytes(["alpha beta gamma delta", "epsilon zeta eta theta", "iota kappa lambda mu"]))
    doc = DoclingDocument(name="d")
    for n in (1, 2, 3):
        doc.add_page(page_no=n, size=Size(width=100, height=100))

    def prov(n, text):
        return ProvenanceItem(page_no=n, bbox=BoundingBox(l=0, t=0, r=1, b=1), charspan=(0, len(text)))

    merged = "alpha beta gamma delta iota kappa lambda mu"
    item = doc.add_text(label=DocItemLabel.TEXT, text=merged, prov=prov(1, merged))
    item.prov.append(prov(3, merged))
    doc.add_text(label=DocItemLabel.TEXT, text="epsilon zeta eta theta", prov=prov(2, "epsilon zeta eta theta"))

    assert page_check.spilled_text(doc) == {3: merged}
    check = page_check.check_document(pdf, doc)
    assert check.status == "checked" and check.findings == []
    assert page_check.check_pages(list(check.source.values()), check.markdown) != [], "page 3 reads as lost without it"
